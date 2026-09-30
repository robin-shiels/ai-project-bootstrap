#!/usr/bin/env python3
"""Summarize workbook structure, formulas, and cached error values as JSON."""

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--max-cells", type=int, default=50000)
    args = parser.parse_args()
    if args.max_cells < 1:
        parser.error("--max-cells must be positive")

    source = load_workbook(args.workbook, read_only=True, data_only=False)
    cached = load_workbook(args.workbook, read_only=True, data_only=True)
    result = {"file": str(args.workbook), "sheets": []}

    for sheet in source:
        values = cached[sheet.title]
        count = 0
        formula_count = 0
        error_count = 0
        formulas = []
        errors = []
        truncated = False
        for row in sheet.iter_rows():
            for cell in row:
                count += 1
                if count > args.max_cells:
                    truncated = True
                    break
                if cell.data_type == "f":
                    formula_count += 1
                    if len(formulas) < 20:
                        formulas.append({"cell": cell.coordinate, "formula": cell.value,
                                         "cached": values[cell.coordinate].value})
                if values[cell.coordinate].data_type == "e":
                    error_count += 1
                    if len(errors) < 20:
                        errors.append({"cell": cell.coordinate, "value": values[cell.coordinate].value})
            if truncated:
                break
        result["sheets"].append({"name": sheet.title, "rows": sheet.max_row,
                                 "columns": sheet.max_column, "cells_scanned": min(count, args.max_cells),
                                 "truncated": truncated, "formulas_scanned": formula_count,
                                 "cached_errors_scanned": error_count,
                                 "formula_samples": formulas, "cached_error_samples": errors})

    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
